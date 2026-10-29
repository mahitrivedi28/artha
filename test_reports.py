import csv
import datetime as dt
import io

from openpyxl import load_workbook
from pypdf import PdfReader

from tests.helpers import auth_headers, register_verify_login


async def _seed_month(client, token):
    cat_resp = await client.get("/api/v1/categories?type=expense", headers=auth_headers(token))
    cat_id = cat_resp.json()["data"][0]["id"]
    today = dt.date.today().replace(day=15)
    await client.post(
        "/api/v1/expenses",
        json={"amount": "500.00", "category_id": cat_id, "merchant": "Test Merchant", "expense_date": today.isoformat()},
        headers=auth_headers(token),
    )
    await client.post(
        "/api/v1/income", json={"amount": "2000.00", "income_date": today.isoformat()}, headers=auth_headers(token)
    )
    return today


async def test_monthly_report_json_matches_real_totals(client, capsys):
    token = await register_verify_login(client, capsys, email="rep1@example.com")
    today = await _seed_month(client, token)

    resp = await client.get(f"/api/v1/reports/monthly?year={today.year}&month={today.month}", headers=auth_headers(token))
    assert resp.status_code == 200, resp.text
    data = resp.json()["data"]
    assert data["total_income"] == "2000.00"
    assert data["total_expense"] == "500.00"
    assert data["total_savings"] == "1500.00"
    assert len(data["transactions"]) == 2


async def test_yearly_report_totals_across_months(client, capsys):
    token = await register_verify_login(client, capsys, email="rep2@example.com")
    today = await _seed_month(client, token)

    resp = await client.get(f"/api/v1/reports/yearly?year={today.year}", headers=auth_headers(token))
    assert resp.status_code == 200
    data = resp.json()["data"]
    assert data["total_income"] == "2000.00"
    assert data["total_expense"] == "500.00"
    assert len(data["monthly_breakdown"]) == 12


async def test_pdf_report_is_a_real_readable_pdf_with_correct_numbers(client, capsys):
    token = await register_verify_login(client, capsys, email="rep3@example.com")
    today = await _seed_month(client, token)

    resp = await client.get(
        f"/api/v1/reports/pdf?report_type=monthly&year={today.year}&month={today.month}", headers=auth_headers(token)
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"

    reader = PdfReader(io.BytesIO(resp.content))
    assert len(reader.pages) >= 1
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert "500.00" in text
    assert "2,000.00" in text  # PDF uses comma thousands-separators for readability
    assert "Test Merchant" in text


async def test_excel_report_is_a_real_readable_workbook_with_correct_numbers(client, capsys):
    token = await register_verify_login(client, capsys, email="rep4@example.com")
    today = await _seed_month(client, token)

    resp = await client.get(
        f"/api/v1/reports/excel?report_type=monthly&year={today.year}&month={today.month}", headers=auth_headers(token)
    )
    assert resp.status_code == 200

    wb = load_workbook(io.BytesIO(resp.content))
    assert "Summary" in wb.sheetnames
    assert "Transactions" in wb.sheetnames

    summary = wb["Summary"]
    values = [cell.value for row in summary.iter_rows() for cell in row if cell.value is not None]
    assert 2000.0 in values
    assert 500.0 in values

    txn_sheet = wb["Transactions"]
    merchant_found = any(
        row[2].value == "Test Merchant" for row in txn_sheet.iter_rows(min_row=2) if row[2].value
    )
    assert merchant_found


async def test_csv_report_parses_correctly(client, capsys):
    token = await register_verify_login(client, capsys, email="rep5@example.com")
    today = await _seed_month(client, token)

    resp = await client.get(
        f"/api/v1/reports/csv?report_type=monthly&year={today.year}&month={today.month}", headers=auth_headers(token)
    )
    assert resp.status_code == 200
    text = resp.content.decode("utf-8")
    rows = list(csv.reader(io.StringIO(text)))
    flat = [cell for row in rows for cell in row]
    assert "500.00" in flat or "500.0" in flat
    assert "Test Merchant" in flat


async def test_yearly_pdf_report_generates_successfully(client, capsys):
    token = await register_verify_login(client, capsys, email="rep6@example.com")
    today = await _seed_month(client, token)

    resp = await client.get(f"/api/v1/reports/pdf?report_type=yearly&year={today.year}", headers=auth_headers(token))
    assert resp.status_code == 200
    reader = PdfReader(io.BytesIO(resp.content))
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert str(today.year) in text


async def test_monthly_download_requires_month(client, capsys):
    token = await register_verify_login(client, capsys, email="rep7@example.com")
    resp = await client.get(f"/api/v1/reports/pdf?report_type=monthly&year=2026", headers=auth_headers(token))
    assert resp.status_code == 422


async def test_reports_isolated_per_user(client, capsys):
    token_a = await register_verify_login(client, capsys, email="repa@example.com")
    token_b = await register_verify_login(client, capsys, email="repb@example.com")
    today = await _seed_month(client, token_a)

    resp = await client.get(f"/api/v1/reports/monthly?year={today.year}&month={today.month}", headers=auth_headers(token_b))
    data = resp.json()["data"]
    assert data["total_expense"] == "0.00"
    assert data["transactions"] == []


async def test_reports_require_auth(client):
    resp = await client.get("/api/v1/reports/monthly?year=2026&month=1")
    assert resp.status_code == 401
